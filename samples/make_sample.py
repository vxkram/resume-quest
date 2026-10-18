"""Render sample_resume.txt to sample_resume.pdf and sample_resume.docx.

The person, companies and university in the sample are fictional.
Needs the dev requirements (reportlab, python-docx): python samples/make_sample.py
"""

from pathlib import Path

import docx
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

HERE = Path(__file__).parent
SECTIONS = {"Summary", "Experience", "Projects", "Skills", "Certifications", "Education"}


def lines():
    return (HERE / "sample_resume.txt").read_text().splitlines()


def make_pdf():
    styles = getSampleStyleSheet()
    story = []
    for i, line in enumerate(lines()):
        if not line.strip():
            story.append(Spacer(1, 6))
        elif i == 0:
            story.append(Paragraph(line, styles["Title"]))
        elif line in SECTIONS:
            story.append(Paragraph(line, styles["Heading2"]))
        elif line.startswith("- "):
            story.append(Paragraph(line[2:], styles["Normal"], bulletText="•"))
        else:
            story.append(Paragraph(line.replace("&", "&amp;"), styles["Normal"]))
    SimpleDocTemplate(str(HERE / "sample_resume.pdf"), pagesize=A4,
                      title="Sample résumé (fictional)", author="Sample").build(story)


def make_docx():
    document = docx.Document()
    for i, line in enumerate(lines()):
        if i == 0:
            document.add_heading(line, level=0)
        elif line in SECTIONS:
            document.add_heading(line, level=1)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            document.add_paragraph(line)
    document.save(HERE / "sample_resume.docx")


if __name__ == "__main__":
    make_pdf()
    make_docx()
    print("wrote sample_resume.pdf and sample_resume.docx")
