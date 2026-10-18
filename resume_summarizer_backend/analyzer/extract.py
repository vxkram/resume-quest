"""Turn an uploaded PDF or DOCX into plain text, entirely in memory.

The file type is decided from its bytes, not from the browser's content type or
the file name, and the upload itself is never written to disk.
"""

import io
import re
import zipfile

from django.conf import settings
from rest_framework.exceptions import APIException

MIN_CHARS = 200
MAX_CHARS = 40_000
MAX_DOCX_UNZIPPED = 50 * 1024 * 1024


class BadUpload(APIException):
    status_code = 400
    default_code = "bad_upload"


def detect_type(data):
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/document.xml" in archive.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            pass
    return None


def _pdf_text(data):
    from pdfminer.high_level import extract_text

    try:
        return extract_text(io.BytesIO(data))
    except Exception as exc:  # pdfminer raises a zoo of exception types on bad input
        raise BadUpload("Couldn't read that PDF. Is it encrypted or damaged?") from exc


def _docx_text(data):
    import docx

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(info.file_size for info in archive.infolist()) > MAX_DOCX_UNZIPPED:
            raise BadUpload("That DOCX expands to an unreasonable size.")
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise BadUpload("Couldn't read that DOCX file.") from exc
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def clean(text):
    text = text.replace("\x00", "").replace("\r", "\n").replace("\f", "\n")
    # pdfminer writes glyphs it can't map (usually bullet symbols) as "(cid:127)".
    text = re.sub(r"^[ \t]*\(cid:\d+\)[ \t]*", "• ", text, flags=re.M)
    text = re.sub(r"\(cid:\d+\)", "", text)
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(uploaded):
    """Return (text, kind) for an uploaded file, or raise BadUpload."""
    if uploaded is None:
        raise BadUpload("Attach a résumé as a PDF or DOCX file.")
    if uploaded.size > settings.MAX_UPLOAD_BYTES:
        raise BadUpload("That file is over the 5 MB limit.")
    data = uploaded.read()
    kind = detect_type(data)
    if kind is None:
        raise BadUpload("Only PDF and DOCX résumés are supported.")
    text = clean(_pdf_text(data) if kind == "pdf" else _docx_text(data))
    if len(text) < MIN_CHARS:
        raise BadUpload("Found almost no text in that file. Scanned PDFs aren't supported.")
    if len(text) > MAX_CHARS:
        raise BadUpload("That document is far longer than a résumé.")
    return text, kind
